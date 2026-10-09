"""Read and validate the ONS labour demand workbook (layout documented in docs/03-workbook.md).

Nothing here touches the database. Anything unexpected raises WorkbookError naming the sheet and cell,
because a workbook that is "mostly parsed" is worse than one that isn't parsed at all.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SUPPRESSED = "[x]"
HEADER_ROW = 5

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH_NAMES = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
               "October", "November", "December"]

# ITL1 region codes (the workbook gives names only).
REGIONS = {
    "North East": "TLC", "North West": "TLD", "Yorkshire and The Humber": "TLE", "East Midlands": "TLF",
    "West Midlands": "TLG", "East of England": "TLH", "London": "TLI", "South East": "TLJ",
    "South West": "TLK", "Wales": "TLL", "Scotland": "TLM", "Northern Ireland": "TLN",
}
COUNTRIES = {"England": "E92000001", "Wales": "W92000004", "Scotland": "S92000003", "Northern Ireland": "N92000002"}
UK = ("K02000001", "United Kingdom")
UNKNOWN_LOCATION = "UNKNOWN"
LAD_CODE = re.compile(r"^[EWSN]\d{8}$")


class WorkbookError(Exception):
    """The workbook doesn't match the layout we know how to read."""


@dataclass(frozen=True)
class TableSpec:
    key: str
    title: re.Pattern[str]  # matched against the text after "Table N: " in cell A1
    columns: tuple[str, ...]  # exact header cells before the period columns
    cadence: str  # "M" or "Q"


TABLES = [
    TableSpec("T1", re.compile(r"^New online job adverts in the UK, split by country between"),
              ("Country",), "M"),
    TableSpec("T2", re.compile(r"^New online job adverts in the UK, split by Local Authority District between"),
              ("Region", "ITL2", "Local Authority District", "Local Authority Code"), "M"),
    TableSpec("T3", re.compile(r"^New online job adverts in the UK, split by region and standard occupation classification \(4-digit SOC2020\)"),
              ("Region", "SOC 4 digit code", "SOC 4 digit label"), "M"),
    TableSpec("T4", re.compile(r"^New online job adverts in the UK, split by Local Authority District and standard occupation classification \(3-digit SOC2020\)"),
              ("Region", "ITL2", "Local Authority District", "SOC 3 digit code", "SOC 3 digit label", "Local Authority Code"), "M"),
    TableSpec("T5", re.compile(r"^Quarterly new online job adverts in the UK, split by Local Authority District and standard occupation classification \(4-digit SOC2020\)"),
              ("Region", "ITL2", "Local Authority District", "SOC 4 digit code", "SOC 4 digit label", "Local Authority Code"), "Q"),
]


@dataclass(frozen=True)
class Location:
    code: str
    name: str
    type: str  # uk | country | region | lad | unknown
    parent_code: str | None = None
    itl2_name: str | None = None


@dataclass(frozen=True)
class Occupation:
    soc_code: str
    name: str
    level: int


@dataclass(frozen=True)
class Observation:
    table: str
    occupation: Occupation
    location: Location
    period: str  # "2026-08" or "2026-Q2"
    adverts: int | None  # None = suppressed by ONS ("[x]")


@dataclass(frozen=True)
class ReleaseInfo:
    release_label: str  # "January 2017 - August 2026"
    published_on: date | None
    latest_period: str  # newest monthly period, "2026-08"


ALL_OCCUPATIONS = Occupation("ALL", "All occupations", 0)


def _cell(sheet: str, row: int, col: int) -> str:
    return f"{sheet}!{get_column_letter(col)}{row}"


def parse_period(header: object, cadence: str, where: str) -> str:
    if not isinstance(header, str):
        raise WorkbookError(f"{where}: expected a period heading, found {header!r}")
    if cadence == "M":
        m = re.fullmatch(r"([A-Z][a-z]{2})-(\d{2})", header.strip())
        if not m or m.group(1) not in MONTHS:
            raise WorkbookError(f"{where}: expected a month like 'Aug-26', found {header!r}")
        return f"20{m.group(2)}-{MONTHS.index(m.group(1)) + 1:02d}"
    m = re.fullmatch(r"(\d{4})Q([1-4])", header.strip())
    if not m:
        raise WorkbookError(f"{where}: expected a quarter like '2026Q2', found {header!r}")
    return f"{m.group(1)}-Q{m.group(2)}"


def next_period(period: str) -> str:
    year, rest = period.split("-")
    if rest.startswith("Q"):
        q = int(rest[1:])
        return f"{int(year) + (q == 4)}-Q{q % 4 + 1}"
    month = int(rest)
    return f"{int(year) + (month == 12)}-{month % 12 + 1:02d}"


def parse_value(value: object, where: str) -> int | None:
    if value == SUPPRESSED:
        return None
    if isinstance(value, bool) or not isinstance(value, int):
        raise WorkbookError(f"{where}: expected a whole number or '[x]', found {value!r}")
    if value < 0:
        raise WorkbookError(f"{where}: negative advert count {value}")
    return value


def _region(name: object, where: str) -> Location:
    if name == "Total UK":
        return Location(*UK, "uk")
    if name == "Unknown":
        return Location(UNKNOWN_LOCATION, "Unknown", "unknown")
    if name not in REGIONS:
        raise WorkbookError(f"{where}: unknown region {name!r}")
    return Location(REGIONS[name], name, "region", parent_code=UK[0])


def _lad(region: object, itl2: object, name: object, code: object, where: str) -> Location:
    parent = _region(region, where)
    if code is None and name == "London" and region == "London":
        return parent  # ONS publishes London as one area at region level (Note 2)
    if not isinstance(code, str) or not LAD_CODE.match(code):
        raise WorkbookError(f"{where}: expected a local authority code like 'E08000003', found {code!r}")
    if not isinstance(name, str) or not name.strip():
        raise WorkbookError(f"{where}: missing local authority name")
    return Location(code, name.strip(), "lad", parent_code=parent.code,
                    itl2_name=itl2.strip() if isinstance(itl2, str) else None)


def _occupation(code: object, label: object, digits: int, where: str) -> Occupation:
    if code == "Unknown":
        return Occupation(f"UNKNOWN-{digits}", "Unknown occupation", 0)
    if not isinstance(code, str) or not re.fullmatch(rf"\d{{{digits}}}", code):
        raise WorkbookError(f"{where}: expected a {digits}-digit SOC code as text, found {code!r}")
    if not isinstance(label, str) or not label.strip():
        raise WorkbookError(f"{where}: missing label for SOC {code}")
    return Occupation(code, label.strip(), digits)


def _row_keys(spec: TableSpec, keys: tuple, where: str) -> tuple[Occupation, Location]:
    match spec.key:
        case "T1":
            (country,) = keys
            if country == "Total UK":
                return ALL_OCCUPATIONS, Location(*UK, "uk")
            if country == "Unknown":
                return ALL_OCCUPATIONS, Location(UNKNOWN_LOCATION, "Unknown", "unknown")
            if country not in COUNTRIES:
                raise WorkbookError(f"{where}: unknown country {country!r}")
            return ALL_OCCUPATIONS, Location(COUNTRIES[country], country, "country", parent_code=UK[0])
        case "T2":
            region, itl2, name, code = keys
            return ALL_OCCUPATIONS, _lad(region, itl2, name, code, where)
        case "T3":
            region, code, label = keys
            return _occupation(code, label, 4, where), _region(region, where)
        case "T4" | "T5":
            region, itl2, name, code, label, lad_code = keys
            digits = 3 if spec.key == "T4" else 4
            return _occupation(code, label, digits, where), _lad(region, itl2, name, lad_code, where)
    raise AssertionError(spec.key)


class Workbook:
    def __init__(self, path: Path):
        try:
            self._wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # openpyxl raises a variety of zip/xml errors
            raise WorkbookError(f"Couldn't open {path.name} as an Excel workbook: {e}") from e
        self._sheets = self._find_tables()

    def close(self) -> None:
        self._wb.close()

    def _find_tables(self) -> dict[str, str]:
        """Map table key -> sheet name by title text, so renumbered or renamed sheets still work."""
        found: dict[str, str] = {}
        for ws in self._wb.worksheets:
            a1 = next(ws.iter_rows(min_row=1, max_row=1, max_col=1, values_only=True), (None,))[0]
            if not isinstance(a1, str):
                continue
            title = re.sub(r"^Table \d+:\s*", "", a1)
            for spec in TABLES:
                if spec.title.search(title):
                    if spec.key in found:
                        raise WorkbookError(f"Two sheets look like {spec.key}: {found[spec.key]!r} and {ws.title!r}")
                    found[spec.key] = ws.title
        missing = [s.key for s in TABLES if s.key not in found]
        if missing:
            raise WorkbookError(f"Couldn't find these tables by title: {', '.join(missing)}. "
                                "The ONS layout may have changed; see docs/03-workbook.md.")
        return found

    def release_info(self) -> ReleaseInfo:
        if "Cover sheet" not in self._wb.sheetnames:
            raise WorkbookError("No 'Cover sheet' worksheet")
        lines = [r[0] for r in self._wb["Cover sheet"].iter_rows(max_col=1, values_only=True) if isinstance(r[0], str)]
        title = lines[0] if lines else ""
        if "Labour Demand Volumes" not in title:
            raise WorkbookError(f"Cover sheet!A1 doesn't look like the labour demand release: {title!r}")
        label = title.split(":", 1)[1].strip() if ":" in title else title
        published = None
        for line in lines:
            m = re.search(r"published on (\d{1,2})(?:st|nd|rd|th)? (\w+) (\d{4})", line)
            if m and m.group(2) in MONTH_NAMES:
                published = date(int(m.group(3)), MONTH_NAMES.index(m.group(2)) + 1, int(m.group(1)))
                break
        header = self._header("T1")
        latest = parse_period(header[-1], "M", "Table 1 header")
        return ReleaseInfo(label, published, latest)

    def data_issues(self) -> list[tuple[str, str]]:
        """(heading, text) for each "Data issue - …" section of the cover sheet, in workbook order.

        A section runs until the next data issue or the next short heading line (e.g. "User guide")."""
        if "Cover sheet" not in self._wb.sheetnames:
            raise WorkbookError("No 'Cover sheet' worksheet")
        issues: list[tuple[str, list[str]]] = []
        current: list[str] | None = None
        for (cell,) in self._wb["Cover sheet"].iter_rows(max_col=1, values_only=True):
            if not isinstance(cell, str) or not cell.strip():
                continue
            line = cell.strip()
            if re.match(r"Data issue\s*-", line, re.I):
                current = []
                issues.append((re.sub(r"\s+", " ", line), current))
            elif current is not None and len(line) < 60 and not line.endswith((".", ":")):
                current = None  # next section heading
            elif current is not None:
                current.append(line)
        for heading, text in issues:
            if not text:
                raise WorkbookError(f"Cover sheet: {heading!r} has no text under it")
        return [(h, "\n\n".join(t)) for h, t in issues]

    def _header(self, key: str) -> tuple:
        ws = self._wb[self._sheets[key]]
        return next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))

    def observations(self, spec: TableSpec) -> Iterator[Observation]:
        sheet = self._sheets[spec.key]
        ws = self._wb[sheet]
        header = tuple(self._header(spec.key))
        while header and header[-1] is None:  # trailing empty columns
            header = header[:-1]
        n = len(spec.columns)
        if header[:n] != spec.columns:
            raise WorkbookError(f"{sheet} row {HEADER_ROW}: expected columns {list(spec.columns)}, "
                                f"found {list(header[:n])}")
        periods = [parse_period(h, spec.cadence, _cell(sheet, HEADER_ROW, n + i + 1)) for i, h in enumerate(header[n:])]
        if not periods:
            raise WorkbookError(f"{sheet}: no period columns")
        for i in range(1, len(periods)):
            if periods[i] != next_period(periods[i - 1]):
                raise WorkbookError(f"{_cell(sheet, HEADER_ROW, n + i + 1)}: periods not consecutive "
                                    f"({periods[i - 1]} then {periods[i]})")

        for r, row in enumerate(ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True), start=HEADER_ROW + 1):
            if all(v is None for v in row):
                continue
            keys = tuple(row[:n])
            occupation, location = _row_keys(spec, keys, _cell(sheet, r, 1))
            for i, period in enumerate(periods):
                col = n + i + 1
                value = row[col - 1] if col - 1 < len(row) else None
                yield Observation(spec.key, occupation, location, period, parse_value(value, _cell(sheet, r, col)))
